"""Presentation regressions; generates artifacts only in temporary directories."""
from pathlib import Path
import sys
from contextlib import redirect_stderr
from io import StringIO
import importlib.util
import tempfile
import unittest
from unittest.mock import patch

for dependency in ['pptx', 'reportlab', 'PIL', 'fitz', 'qrcode']:
    if importlib.util.find_spec(dependency) is None:
        raise unittest.SkipTest('Install presentation requirements to run slide tests: '+dependency)
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'docs/workshops/presentations'))
import build
from content import DECKS
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
import fitz


class Workshop4Presentation(unittest.TestCase):
    def test_keyed_outputs_reject_repository_and_symlink_paths(self):
        with tempfile.TemporaryDirectory() as directory:
            key = Path(directory)/'key.txt'
            key.write_text('TEST-ONLY-KEY')
            alias = Path(directory)/'repository-alias'
            alias.symlink_to(build.ROOT, target_is_directory=True)
            for target in [build.ROOT, build.ROOT/'event-output', alias]:
                with self.subTest(target=target), patch.object(build, 'HERE', build.ROOT/'docs/workshops/presentations'), patch.object(build, 'AUDIENCE_KEY', None), patch.object(build, 'make_deck', side_effect=AssertionError('Repository output was not rejected')) as generate, patch.object(sys, 'argv', ['build.py','--workshop','4','--audience-key-file',str(key),'--output-dir',str(target)]):
                    with redirect_stderr(StringIO()), self.assertRaises(SystemExit) as caught:
                        build.main()
                    self.assertEqual(caught.exception.code, 2)
                    generate.assert_not_called()

    def test_architecture_fits_and_remains_editable(self):
        deck = next(deck for deck in DECKS if deck['number'] == 4)
        architecture = next(slide for slide in deck['slides'] if slide['kind'] == 'w4_architecture')
        one_slide = {**deck, 'slides': [architecture]}
        with tempfile.TemporaryDirectory() as directory, patch.object(build, 'HERE', Path(directory)):
            stem, audits = build.make_deck(one_slide)
            self.assertEqual(build.verify_and_preview(one_slide, stem, audits)['checks'], 'passed')
            prs = Presentation(Path(directory)/(stem+'.pptx'))
            self.assertFalse(any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in prs.slides[0].shapes))
            text = '\n'.join(shape.text for shape in prs.slides[0].shapes if shape.has_text_frame)
            for label in ['Gate 1 / AI', 'Gate 2 / MCP', 'Gate 3 / API', 'Independent reviewer', 'Core banking', 'OPA', 'Jaeger']:
                self.assertIn(label, text)

    def test_gateway_primer_precedes_incident_workshop(self):
        deck = next(deck for deck in DECKS if deck['number'] == 4)
        titles = [slide['title'].lower() for slide in deck['slides']]
        self.assertIn('api gateway', titles[0])
        api = next(i for i, title in enumerate(titles) if 'api gateway' in title)
        ai = next(i for i, title in enumerate(titles) if 'ai gateway' in title)
        mcp = next(i for i, title in enumerate(titles) if 'mcp gateway' in title)
        incident = titles.index('the day the agent broke the bank')
        self.assertLess(api, ai)
        self.assertLess(ai, mcp)
        self.assertLess(mcp, incident)
        self.assertTrue(any(slide['kind'] == 'w4_architecture' for slide in deck['slides']))
        self.assertTrue(any(slide['kind'] == 'audience_access' for slide in deck['slides']))
        self.assertTrue(any('printenv W4_REVIEWER_PASSWORD' in (slide['code'] or '') for slide in deck['slides']))

    def test_event_access_slide_contains_scannable_qr_and_key(self):
        deck = next(deck for deck in DECKS if deck['number'] == 4)
        access = next((slide for slide in deck['slides'] if slide['kind'] == 'audience_access'), None)
        self.assertIsNotNone(access, 'Deck needs an audience access slide')
        one_slide = {**deck, 'slides': [access]}
        with tempfile.TemporaryDirectory() as directory, patch.object(build, 'HERE', Path(directory)), patch.object(build, 'AUDIENCE_KEY', 'TEST-EVENT-KEY', create=True):
            stem, audits = build.make_deck(one_slide)
            result = build.verify_and_preview(one_slide, stem, audits)
            self.assertEqual(result['checks'], 'passed')
            prs = Presentation(Path(directory) / (stem+'.pptx'))
            pictures = [shape for shape in prs.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
            self.assertEqual(len(pictures), 1)
            with fitz.open(Path(directory) / (stem+'.pdf')) as pdf:
                self.assertIn('TEST-EVENT-KEY', pdf[0].get_text())
            if importlib.util.find_spec('cv2') is None:
                self.skipTest('Install optional opencv-python-headless for QR decoding')
            import cv2
            import numpy as np
            decoded, _, _ = cv2.QRCodeDetector().detectAndDecode(cv2.imdecode(np.frombuffer(pictures[0].image.blob, dtype=np.uint8), cv2.IMREAD_COLOR))
            self.assertEqual(decoded, 'https://w4.ravirai.in/workshop-4')
            self.assertNotIn('TEST-EVENT-KEY', decoded)


if __name__ == '__main__':
    unittest.main()
