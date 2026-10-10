#!/usr/bin/env python3
"""Generate editable PowerPoint slides, PDF handouts and presenter notes.

Run with the dependencies in requirements.txt. Both slide formats use the same
layout and explicit line wrapping. PDF copies are rendered independently, not
exported by PowerPoint; compare in your presentation app after editing a deck.
"""
from pathlib import Path
import argparse
from io import BytesIO
import json
import math
import zipfile

from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.utils import ImageReader
import fitz
from PIL import Image, ImageDraw
import qrcode

from content import DECKS

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
AUDIENCE_URL = 'https://w4.ravirai.in/workshop-4'
AUDIENCE_KEY = None
AUDIENCE_EXPIRY = '12:00 IST · 10 October 2026'
W, H = 960, 540
NAVY, PAPER, INK = "102B36", "F5F7F4", "163640"
MUTED, TEAL, WHITE, BORDER = "536B73", "007F79", "FFFFFF", "D8E3DF"
ACCENTS = ["37D7BD", "F5BC69", "A9B8FF", "FF9982"]
FONTDIR = Path("/usr/share/fonts/truetype/dejavu")
for face, filename in [("Deck", "DejaVuSans.ttf"), ("DeckBold", "DejaVuSans-Bold.ttf"),
                       ("DeckMono", "DejaVuSansMono.ttf")]:
    pdfmetrics.registerFont(TTFont(face, str(FONTDIR / filename)))


def wrapped(text, face, size, width):
    lines = []
    for line in text.split("\n"):
        if not line:
            lines.append("")
            continue
        current = ""
        for word in line.split(" "):
            test = current + (" " if current else "") + word
            if pdfmetrics.stringWidth(test, face, size) <= width:
                current = test
            else:
                if current:
                    lines.append(current)
                current = word
                if pdfmetrics.stringWidth(word, face, size) > width:
                    # Long identifiers wrap without dropping or changing characters.
                    current = ""
                    for char in word:
                        if current and pdfmetrics.stringWidth(current + char, face, size) > width:
                            lines.append(current)
                            current = char
                        else:
                            current += char
        lines.append(current)
    return lines


class Painter:
    def __init__(self, slide, pdf, audit):
        self.slide, self.pdf, self.audit = slide, pdf, audit

    def rect(self, x, y, w, h, color, radius=False):
        shape = self.slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
            Pt(x), Pt(y), Pt(w), Pt(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(color)
        shape.line.fill.background()
        if radius:
            shape.adjustments[0] = 0.06
        self.pdf.setFillColorRGB(*[int(color[i:i+2], 16)/255 for i in (0, 2, 4)])
        if radius:
            self.pdf.roundRect(x, H-y-h, w, h, 7, stroke=0, fill=1)
        else:
            self.pdf.rect(x, H-y-h, w, h, stroke=0, fill=1)

    def text(self, text, x, y, w, h, size=20, color=INK, bold=False,
             mono=False, minimum=14, leading=1.25):
        face = "DeckMono" if mono else "DeckBold" if bold else "Deck"
        fontname = "DejaVu Sans Mono" if mono else "DejaVu Sans"
        actual = size
        while actual >= minimum:
            lines = text.split("\n") if mono else wrapped(text, face, actual, w-4)
            lineheight = actual * leading
            widest_token = max((pdfmetrics.stringWidth(token, face, actual)
                                for token in text.split()), default=0)
            widest_line = max((pdfmetrics.stringWidth(line, face, actual)
                               for line in lines), default=0)
            if (len(lines) * lineheight <= h-4 and widest_line <= w-4
                    and (widest_token <= w-4 or actual <= minimum)):
                break
            actual -= 0.5
        if actual < minimum:
            raise ValueError(f"Text overflow: {text!r}, {w}x{h}")
        self.audit.append(dict(text=text, font_size=actual, x=x, y=y, w=w, h=h,
                               lines=len(lines), line_height=lineheight))
        box = self.slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
        tf = box.text_frame
        tf.clear()
        tf.word_wrap = False
        tf.margin_left = tf.margin_right = Pt(0)
        tf.margin_top = tf.margin_bottom = Pt(0)
        tf.vertical_anchor = MSO_ANCHOR.TOP
        for index, line in enumerate(lines):
            p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
            p.text = line
            p.font.name = fontname
            p.font.size = Pt(actual)
            p.font.bold = bold
            p.font.color.rgb = RGBColor.from_string(color)
            p.space_before = p.space_after = Pt(0)
            p.line_spacing = Pt(lineheight)
        self.pdf.setFillColorRGB(*[int(color[i:i+2], 16)/255 for i in (0, 2, 4)])
        self.pdf.setFont(face, actual)
        ascent = pdfmetrics.getAscent(face) * actual / 1000
        for index, line in enumerate(lines):
            self.pdf.drawString(x, H-y-ascent-index*lineheight, line)

    def image(self, png, x, y, w, h):
        self.slide.shapes.add_picture(BytesIO(png), Pt(x), Pt(y), Pt(w), Pt(h))
        self.pdf.drawImage(ImageReader(BytesIO(png)), x, H-y-h, w, h)


def audience_qr():
    qr = qrcode.QRCode(error_correction=qrcode.constants.ERROR_CORRECT_Q, box_size=12, border=4)
    qr.add_data(AUDIENCE_URL)
    qr.make(fit=True)
    output = BytesIO()
    qr.make_image(fill_color='black', back_color='white').save(output, format='PNG')
    return output.getvalue()


def footer(p, deck, page, total, dark=False):
    color = "AAC0C6" if dark else MUTED
    p.text(f"FLO BANK  /  WORKSHOP {deck['number']}  /  FICTIONAL LAB", 42, 515, 700, 18,
           size=9, minimum=9, color=color)
    p.text(f"{page:02d} / {total:02d}", 847, 514, 75, 20, size=10, minimum=10, color=color)


def render(p, d, s, page):
    accent = ACCENTS[d["number"]-1]
    dark = s["kind"] in ("cover", "statement")
    p.rect(0, 0, W, H, NAVY if dark else PAPER)
    p.rect(42, 30, 36, 4, accent if dark else TEAL)
    p.text(f"WORKSHOP {d['number']}  /  {s['time'].upper()}", 92, 22, 810, 24,
           size=11, minimum=10, color=accent if dark else TEAL, bold=True)
    if dark:
        if s["kind"] == "cover":
            p.rect(761, 85, 156, 316, "193D48", radius=True)
            p.text(f"0{d['number']}", 782, 118, 126, 140, size=82, minimum=72, color=accent, bold=True)
            p.text(f"{d['duration']}\nMINUTES", 782, 276, 123, 100, size=25, minimum=22, color=WHITE)
            p.text(s["title"], 42, 108, 682, 202, size=46, minimum=36, color=WHITE, bold=True, leading=1.15)
            p.text(s["lead"], 44, 327, 660, 102, size=22, minimum=19, color="C1D5D8")
        else:
            p.text(s["title"], 42, 107, 870, 178, size=44, minimum=32, color=WHITE, bold=True, leading=1.15)
            p.text(s["lead"], 44, 300, 851, 121, size=26, minimum=21, color="C1D5D8")
        p.text(s["takeaway"], 44, 459, 863, 45, size=17, minimum=14, color=accent, bold=True)
        footer(p, d, page, len(d["slides"]), True)
        return

    p.text(s["title"], 42, 65, 876, 80, size=30, minimum=25, bold=True, leading=1.14)
    top = 169
    if s["lead"]:
        p.text(s["lead"], 43, 145, 871, 58, size=18, minimum=16, color=MUTED)
        top = 215
    bottom = 444
    if s["kind"] == "audience_access":
        p.image(audience_qr(), 65, 170, 260, 260)
        p.text('OPEN ON YOUR PHONE', 365, 172, 520, 28, size=14, color=TEAL, bold=True)
        p.text(AUDIENCE_URL, 365, 209, 520, 60, size=24, minimum=21, bold=True)
        p.text('AUDIENCE ACCESS CODE', 365, 285, 520, 26, size=14, color=TEAL, bold=True)
        p.text(AUDIENCE_KEY or 'Provided live by facilitator', 365, 318, 520, 52, size=30, minimum=23, bold=True)
        p.text('Recorded evidence only · no payment execution', 365, 385, 520, 27, size=16, minimum=14)
        p.text('Access closes: '+AUDIENCE_EXPIRY, 365, 421, 520, 28, size=13, minimum=12, color=MUTED)
    elif s["kind"] == "w4_architecture":
        def node(x, y, title, body, width=180, height=82, color=WHITE):
            p.rect(x, y, width, height, color, radius=True)
            compact = height < 60
            p.text(title, x+10, y+7, width-20, 23 if compact else 29, size=14, minimum=12, bold=True, color=TEAL)
            offset = 28 if compact else 39
            p.text(body, x+10, y+offset, width-20, height-offset-2, size=11 if compact else 12, minimum=10, leading=1.15)

        def arrow(x, y, direction='→'):
            p.text(direction, x, y, 40, 28, size=22, minimum=18, color=TEAL)

        node(42, 157, 'Agent backend', 'Orchestrate requests;\nvalidate model proposals')
        arrow(236, 180)
        node(272, 157, 'Gate 1 / AI', 'Inference budget;\nreserve before dispatch')
        arrow(466, 180)
        node(502, 157, 'Model provider', 'Returns a proposal\nor refusal')
        node(742, 157, 'Independent reviewer', 'Approve exact proposal;\nseparate session', width=176, color='E0ECE6')
        p.text('Model output returns to the agent; the agent requests tools below.', 42, 244, 640, 24, size=12, minimum=11, color=MUTED)
        arrow(807, 244, '↓')
        node(42, 279, 'Tool request', 'Named capability +\nactual payment arguments')
        arrow(236, 302)
        node(272, 279, 'Gate 2 / MCP', 'Caller + tool arguments;\nOPA policy decision')
        arrow(466, 302)
        node(502, 279, 'Gate 3 / API', 'APISIX route + API\naudience and scope checks')
        arrow(701, 302)
        node(742, 279, 'Core banking', 'Approval binding;\nidempotent settlement', width=176)
        arrow(345, 361, '↕')
        arrow(807, 361, '↕')
        node(272, 393, 'OPA', 'Role / amount / beneficiary', height=53)
        node(502, 393, 'Jaeger', 'Observed distributed traces', height=53)
        node(742, 393, 'Durable records', 'Tasks / proposals / payments', width=176, height=53)
        p.text('Three logical gates share APISIX.\nThe isolated vulnerable ledger\nis a separate demonstration.', 42, 381, 210, 68, size=12, minimum=10, color=MUTED, leading=1.2)
    elif s["kind"] == "w2_architecture":
        def node(x, y, w, h, title, body, color=WHITE):
            p.rect(x, y, w, h, color, radius=True)
            p.text(title, x+10, y+9, w-20, 25, size=14, minimum=12, bold=True, color=TEAL)
            p.text(body, x+10, y+36, w-20, h-40, size=12, minimum=10, leading=1.15)

        def arrow(x, y, w=30, direction="→"):
            p.text(direction, x, y, w, 28, size=20, minimum=16, color=TEAL)

        node(42, 158, 185, 82, "Flo / Studio backend", "Read case via MCP first;\nvalidate model proposal")
        arrow(234, 180)
        node(267, 158, 190, 82, "Gate 1 / AI adapter", "Inference token budget\nand output limits")
        arrow(465, 180)
        node(499, 158, 185, 82, "Hosted LLM", "Live review: returns a\nproposal or refusal")
        p.text("Proposal returns to Flo", 699, 168, 215, 26, size=12, minimum=11, color=MUTED)
        p.text("then enters MCP below ↓", 699, 199, 215, 28, size=12, minimum=11, color=MUTED)

        node(42, 270, 185, 87, "Payment proposal", "Live: validated model call\nReplay: fixed request", "E0ECE6")
        arrow(234, 295)
        node(267, 270, 190, 87, "Gate 2 / MCP adapter", "Signed identity + arguments\nOPA decision; fail closed")
        arrow(465, 295)
        node(499, 270, 185, 87, "Gate 3 / API gateway", "API-key authentication\nForward authorized calls")
        arrow(692, 295)
        node(727, 270, 190, 87, "Core Banking", "JWT + business rules\nApprovals and ledger")

        arrow(349, 357, direction="↕")
        node(267, 387, 190, 58, "OPA", "Role / amount / beneficiary")
        arrow(808, 357, direction="↕")
        node(727, 387, 190, 58, "PostgreSQL", "Banking records")
        node(499, 387, 185, 58, "Jaeger", "Exported distributed traces")
        p.text("APISIX hosts Gates 1–3.\nDenied payments stop at Gate 2.\nPending approval: no debit.",
               42, 382, 205, 65, size=12, minimum=10, color=MUTED, leading=1.2)
    elif s["cards"]:
        cards = s["cards"]
        cw = (876 - 20 * (len(cards)-1)) / len(cards)
        for i, (heading, body) in enumerate(cards):
            x = 42 + i * (cw+20)
            p.rect(x, top, cw, bottom-top, WHITE, radius=True)
            p.rect(x+19, top+20, 30, 4, TEAL)
            p.text(heading, x+19, top+42, cw-38, 55, size=16, minimum=13, bold=True, color=TEAL)
            p.text(body, x+19, top+104, cw-38, bottom-top-117, size=21, minimum=16.5)
    elif s["steps"]:
        steps = s["steps"]
        cw = (876-22*(len(steps)-1)) / len(steps)
        for i, (heading, body) in enumerate(steps):
            x = 42+i*(cw+22)
            p.rect(x, top+24, cw, bottom-top-33, WHITE, radius=True)
            p.text(f"{i+1:02d}", x+15, top+40, cw-30, 38, size=22, color=TEAL, bold=True)
            p.text(heading, x+15, top+90, cw-30, 65, size=18, minimum=15, bold=True)
            p.text(body, x+15, top+157, cw-30, bottom-top-176, size=17, minimum=14.5)
            if i<len(steps)-1:
                p.text("→", x+cw+3, top+111, 20, 26, size=18, minimum=16, color=TEAL)
    elif s["code"]:
        p.rect(42, top, 876, bottom-top, NAVY, radius=True)
        p.text("JSON-RPC  /  SANITIZED REQUEST BODY" if s['kind'] == 'request_payload' else "TERMINAL  /  FROM REPOSITORY ROOT", 61, top+15, 820, 25,
               size=10, minimum=10, color=accent, bold=True)
        p.text(s["code"], 61, top+54, 837, bottom-top-69,
               size=17, minimum=13, color=WHITE, mono=True, leading=1.4)
    elif s["rows"]:
        headers, rows = s["headers"], s["rows"]
        n = len(headers)
        weights = {2:[0.47,0.53],3:[0.30,0.36,0.34],4:[0.24,0.21,0.27,0.28]}[n]
        if headers[0] == "Time":
            weights = [0.12,0.44,0.44]
        widths = [876*v for v in weights]
        header_h = 40
        rh = (bottom-top-header_h)/len(rows)
        p.rect(42, top, 876, header_h, NAVY)
        x=42
        for heading, width in zip(headers, widths):
            p.text(heading, x+12, top+10, width-24, 29, size=14, minimum=11, color=WHITE, bold=True)
            x+=width
        for j, row in enumerate(rows):
            y=top+header_h+j*rh
            p.rect(42,y,876,rh,WHITE if j%2==0 else "E9EFEB")
            x=42
            for cell, width in zip(row,widths):
                p.text(cell,x+12,y+7,width-24,rh-9,size=17,minimum=12.5,leading=1.16)
                x+=width
    p.rect(42, 461, 876, 43, "E0ECE6", radius=True)
    p.text(s["takeaway"] or "Follow the worksheet; record the actual evidence from your run.",
           55, 470, 848, 31, size=15, minimum=12.5, bold=True)
    footer(p, d, page, len(d["slides"]))


def make_deck(d):
    stem = f"flo-bank-workshop-{d['number']}"
    ppt_path, pdf_path = HERE / f"{stem}.pptx", HERE / f"{stem}.pdf"
    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(W), Pt(H)
    prs.core_properties.title = d["title"]
    prs.core_properties.subject = d["subtitle"]
    prs.core_properties.author = "Flo Bank Workshop Series"
    prs.core_properties.keywords = f"Flo Bank, workshop {d['number']}, fictional lab"
    pdf = canvas.Canvas(str(pdf_path), pagesize=(W,H))
    pdf.setTitle(d["title"])
    pdf.setAuthor("Flo Bank Workshop Series")
    notes = [f"# Workshop {d['number']}: {d['title']}", "", d["subtitle"], "",
             f"Duration: {d['duration']} minutes. Last slide is presenter reference outside the timed agenda.", "",
             "Expected outcomes in these slides are instructional targets, not fresh execution results.", ""]
    if d['number'] == 4:
        notes += ['Delivery: 15-minute gateway primer followed by the 135-minute lab, including its seven-minute break.', '']
    audits=[]
    for i,s in enumerate(d["slides"],1):
        sl=prs.slides.add_slide(prs.slide_layouts[6])
        audit=[]
        render(Painter(sl,pdf,audit),d,s,i)
        source_text="\n".join(d["sources"])
        sl.notes_slide.notes_text_frame.text=(f"SLIDE {i} / {s['time']}\n{s['notes']}\n\n"
                                              f"SOURCE FILES (repository root):\n{source_text}")
        notes += [f"## {i:02d}. {s['title']}", "", f"**Timing:** {s['time']}", "",
                  s["notes"], ""]
        if s["code"]:
            notes += ["```bash",s["code"],"```",""]
        notes += ["Sources: " + ", ".join(f"[{Path(x).name}](../../../{x})" for x in d["sources"]), ""]
        audits.append(audit)
        pdf.showPage()
    prs.save(ppt_path)
    pdf.save()
    (HERE/f"{stem}-notes.md").write_text("\n".join(notes),encoding="utf-8")
    return stem,audits


def verify_and_preview(d, stem, audits):
    ppt_path, pdf_path = HERE/f"{stem}.pptx", HERE/f"{stem}.pdf"
    prs=Presentation(ppt_path)
    doc=fitz.open(pdf_path)
    expected=len(d["slides"])
    assert len(prs.slides)==len(doc)==expected
    with zipfile.ZipFile(ppt_path) as z:
        assert z.testzip() is None
    for source in d["sources"]:
        assert (ROOT/source).is_file(), source
    for i,(sl,page,audit) in enumerate(zip(prs.slides,doc,audits),1):
        assert sl.has_notes_slide and "SOURCE FILES" in sl.notes_slide.notes_text_frame.text
        assert page.get_text().strip(), f"Empty PDF page {i}"
        for shape in sl.shapes:
            assert shape.left>=0 and shape.top>=0
            assert shape.left+shape.width<=prs.slide_width+Pt(1)
            assert shape.top+shape.height<=prs.slide_height+Pt(1)
        for item in audit:
            assert item["lines"]*item["line_height"]<=item["h"]-4+0.01
        for block in page.get_text("dict")["blocks"]:
            if block["type"]==0:
                x0,y0,x1,y1=block["bbox"]
                assert x0>=0 and y0>=0 and x1<=W and y1<=H, (i,block)
    cols,tw,th=3,480,270
    sheet=Image.new("RGB",(cols*tw, math.ceil(expected/cols)*(th+28)),"#D8E3DF")
    draw=ImageDraw.Draw(sheet)
    for i,page in enumerate(doc):
        pix=page.get_pixmap(matrix=fitz.Matrix(0.5,0.5),alpha=False)
        im=Image.frombytes("RGB",(pix.width,pix.height),pix.samples)
        x=(i%cols)*tw;y=(i//cols)*(th+28)
        sheet.paste(im,(x,y))
        draw.text((x+10,y+th+6),f"{i+1:02d} | {d['slides'][i]['time']}",fill="#163640")
    preview_dir=HERE/"previews"
    preview_dir.mkdir(exist_ok=True)
    sheet.save(preview_dir/f"{stem}-contact-sheet.png")
    doc[0].get_pixmap(matrix=fitz.Matrix(1,1),alpha=False).save(str(preview_dir/f"{stem}-cover.png"))
    # Inspect selected layouts at full size in addition to the contact sheet.
    for idx in (min(5,expected-1), min(10,expected-1)):
        doc[idx].get_pixmap(matrix=fitz.Matrix(1,1),alpha=False).save(str(preview_dir/f"{stem}-slide-{idx+1:02d}.png"))
    for idx, slide in enumerate(d['slides']):
        if slide['kind'] in ('w4_architecture', 'audience_access'):
            label = 'architecture' if slide['kind'] == 'w4_architecture' else 'audience-access'
            doc[idx].get_pixmap(matrix=fitz.Matrix(1.5,1.5),alpha=False).save(str(preview_dir/f'{stem}-{label}.png'))
    doc.close()
    return dict(workshop=d["number"], slides=expected, duration_minutes=d["duration"],
                pptx=ppt_path.name,pdf=pdf_path.name, checks="passed",
                min_font_size=min(a["font_size"] for slide in audits for a in slide),
                validation=["PPTX reopened", "ZIP integrity", "PDF page count and text bounds",
                            "Slide shape bounds", "Text fit", "Presenter notes", "Source paths"],
                rendering_note="PDF uses the same layout, independently rendered; no PowerPoint/LibreOffice renderer available.")


def main():
    global HERE, AUDIENCE_KEY, AUDIENCE_EXPIRY
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--workshop', type=int, choices=[1,2,3,4])
    parser.add_argument('--output-dir', type=Path)
    parser.add_argument('--audience-key-file', type=Path)
    parser.add_argument('--audience-expiry', default=AUDIENCE_EXPIRY)
    args = parser.parse_args()
    if args.audience_key_file:
        if not args.output_dir or args.workshop != 4:
            parser.error('Event keys require --workshop 4 and a separate --output-dir; keep keyed decks out of repository outputs.')
        target = args.output_dir.resolve()
        if target.is_relative_to(ROOT):
            parser.error('Generate the event deck outside the repository to avoid publishing its access code.')
        AUDIENCE_KEY = args.audience_key_file.read_text().strip()
        if not AUDIENCE_KEY:
            parser.error('Audience key file is empty')
    if args.output_dir:
        HERE = args.output_dir.resolve()
        HERE.mkdir(parents=True, exist_ok=True)
    AUDIENCE_EXPIRY = args.audience_expiry
    selected = [d for d in DECKS if not args.workshop or d['number'] == args.workshop]
    results=[]
    for d in selected:
        stem,audits=make_deck(d)
        result=verify_and_preview(d,stem,audits)
        results.append(result)
        print(f"W{d['number']}: {result['slides']} slides; PPTX, PDF, notes and previews generated; checks passed.")
    validation_path = HERE/'validation.json'
    if args.workshop and validation_path.exists():
        prior = json.loads(validation_path.read_text())
        results = sorted([r for r in prior if r['workshop'] != args.workshop]+results, key=lambda r:r['workshop'])
    validation_path.write_text(json.dumps(results,indent=2)+"\n")
    if any(d['number'] == 4 for d in selected):
        (HERE/'w4-audience-qr.png').write_bytes(audience_qr())
    bundle=HERE/(f'flo-bank-workshop-{args.workshop}-package.zip' if args.workshop else 'flo-bank-all-workshops.zip')
    with zipfile.ZipFile(bundle,"w",compression=zipfile.ZIP_DEFLATED) as z:
        for path in sorted(HERE.iterdir()):
            include = not args.workshop or path.name.startswith(f'flo-bank-workshop-{args.workshop}') or path.name == 'w4-audience-qr.png'
            if include and path.suffix in (".pptx",".pdf",".md",".py",".txt",".json",".png"):
                z.write(path,path.name)
    print(f"Bundle: {bundle}")


if __name__=="__main__":
    main()
